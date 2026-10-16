import io
import zipfile
import uuid
import pytest
from datetime import datetime, date, timedelta, timezone
from unittest.mock import AsyncMock, patch, MagicMock
from httpx import AsyncClient, ASGITransport
from sqlalchemy.ext.asyncio import async_sessionmaker, AsyncSession
from docx import Document

from app.main import app
from app.config import settings
from app.core.security import get_password_hash, create_access_token
from app.models.organization import Organization
from app.models.user import User, OrganizationMember
from app.models.contract import (
    Contract,
    ContractVersion,
    ContractRedline
)
from app.models.obligation import Obligation
from app.models.notification import Notification
from app.models.audit import AuditLog
from app.services.notifications.scheduler import NotificationScheduler, NOTIFICATION_ADVISORY_LOCK_KEY
from app.services.notifications.worker import NotificationWorker
from app.services.export.docx_export import generate_contract_revisions_docx, sanitize_filename


async def _setup_tenant(test_engine, name_prefix: str = "Phase2", role: str = "OWNER"):
    suffix = uuid.uuid4().hex[:8]
    TestSessionLocal = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)
    async with TestSessionLocal() as session:
        org = Organization(name=f"{name_prefix} Org {suffix}", slug=f"{name_prefix.lower()}-org-{suffix}")
        session.add(org)
        await session.flush()

        user = User(
            email=f"{name_prefix.lower()}_{suffix}@example.com",
            password_hash=get_password_hash("Password123!"),
            full_name=f"{name_prefix} User {suffix}"
        )
        session.add(user)
        await session.flush()

        member = OrganizationMember(
            organization_id=org.id,
            user_id=user.id,
            role=role,
            status="ACTIVE"
        )
        session.add(member)
        await session.commit()

        token = create_access_token(subject=user.id, org_id=org.id, role=role)
        return {
            "org_id": org.id,
            "user_id": user.id,
            "email": user.email,
            "token": token,
            "headers": {"Authorization": f"Bearer {token}"},
            "suffix": suffix
        }


# ==============================================================================
# Task A Tests: Notification Scheduler & Background Coordination
# ==============================================================================

@pytest.mark.asyncio
async def test_scheduler_lifecycle():
    """Verify scheduler startup, status reporting, and clean shutdown without task leakage."""
    scheduler = NotificationScheduler()
    assert not scheduler.is_running
    assert scheduler.task is None

    # Start the scheduler
    task = scheduler.start()
    assert scheduler.is_running
    assert scheduler.task is not None
    assert scheduler.task == task

    # Calling start again while already running returns existing task
    redundant_task = scheduler.start()
    assert redundant_task == task

    # Shutdown cleanly
    await scheduler.stop()
    assert not scheduler.is_running
    assert scheduler.task is None


@pytest.mark.asyncio
async def test_scheduler_in_process_concurrency_lock():
    """Verify that multiple concurrent calls to execute_cycle are blocked by the in-process lock."""
    scheduler = NotificationScheduler()

    # Manually lock the in-process lock to simulate an ongoing evaluation cycle
    await scheduler._lock.acquire()
    try:
        res = await scheduler.execute_cycle()
        assert res["status"] == "skipped"
        assert res["reason"] == "local_lock_busy"
    finally:
        scheduler._lock.release()


@pytest.mark.asyncio
async def test_scheduler_worker_failure_resilience():
    """Verify that worker runtime exceptions are caught, logged safely, and do not crash the scheduler."""
    scheduler = NotificationScheduler()

    with patch.object(
        NotificationWorker,
        "evaluate_and_dispatch",
        new_callable=AsyncMock,
        side_effect=RuntimeError("Simulated intermittent database disconnection")
    ):
        res = await scheduler.execute_cycle()
        assert res["status"] == "error"
        assert "Simulated intermittent" in res["error"]

    # Verify lock was released cleanly despite failure
    assert not scheduler._lock.locked()


@pytest.mark.asyncio
async def test_scheduler_postgresql_advisory_lock_branch():
    """Verify PostgreSQL advisory lock acquisition, busy-skip, and release logic."""
    scheduler = NotificationScheduler()

    # Mock settings to simulate PostgreSQL database URL
    with patch.object(settings, "DATABASE_URL", "postgresql+asyncpg://user:pass@localhost/db"):
        # Case 1: pg_try_advisory_lock returns False (another API instance holds the lock)
        mock_result_busy = MagicMock()
        mock_result_busy.scalar.return_value = False

        mock_session_busy = AsyncMock()
        mock_session_busy.execute = AsyncMock(return_value=mock_result_busy)

        ctx_busy = MagicMock()
        ctx_busy.__aenter__ = AsyncMock(return_value=mock_session_busy)
        ctx_busy.__aexit__ = AsyncMock(return_value=None)

        mock_engine_busy = MagicMock()
        mock_engine_busy.connect.return_value = ctx_busy

        with patch("app.services.notifications.scheduler.engine", mock_engine_busy):
            with patch.object(NotificationWorker, "evaluate_and_dispatch", new_callable=AsyncMock) as mock_worker:
                res = await scheduler.execute_cycle()
                assert res["status"] == "skipped"
                assert res["reason"] == "postgres_advisory_lock_busy"
                mock_worker.assert_not_called()

        # Case 2: pg_try_advisory_lock returns True (lock acquired)
        mock_result_free = MagicMock()
        mock_result_free.scalar.return_value = True

        mock_conn_free = AsyncMock()
        mock_conn_free.execute = AsyncMock(return_value=mock_result_free)

        ctx_free = MagicMock()
        ctx_free.__aenter__ = AsyncMock(return_value=mock_conn_free)
        ctx_free.__aexit__ = AsyncMock(return_value=None)

        mock_engine_free = MagicMock()
        mock_engine_free.connect.return_value = ctx_free

        with patch("app.services.notifications.scheduler.engine", mock_engine_free):
            with patch.object(
                NotificationWorker,
                "evaluate_and_dispatch",
                new_callable=AsyncMock,
                return_value={"evaluated_contracts": 1, "notifications_created": 1, "skipped_dedup": 0}
            ):
                res = await scheduler.execute_cycle()
                assert res["status"] == "completed"
                # Check that pg_try_advisory_lock and pg_advisory_unlock were both called on lock_conn
                assert mock_conn_free.execute.call_count >= 2


@pytest.mark.asyncio
async def test_scheduler_sqlite_fallback_branch():
    """Verify SQLite fallback logic when running in local development mode without distributed locks."""
    scheduler = NotificationScheduler()

    # Ensure SQLite mode is configured
    with patch.object(settings, "DATABASE_URL", "sqlite+aiosqlite:///contractly.db"):
        mock_session = AsyncMock()
        ctx = MagicMock()
        ctx.__aenter__ = AsyncMock(return_value=mock_session)
        ctx.__aexit__ = AsyncMock(return_value=None)
        with patch("app.services.notifications.scheduler.AsyncSessionLocal", return_value=ctx):
            with patch.object(
                NotificationWorker,
                "evaluate_and_dispatch",
                new_callable=AsyncMock,
                return_value={"evaluated_contracts": 2, "notifications_created": 0, "skipped_dedup": 0}
            ):
                res = await scheduler.execute_cycle()
                assert res["status"] == "completed"
                # In SQLite mode, no pg_try_advisory_lock query is run
                mock_session.execute.assert_not_called()


@pytest.mark.asyncio
async def test_notification_evaluation_deduplication(test_engine):
    """Verify that multiple consecutive worker evaluation runs do not generate duplicate notifications."""
    tenant = await _setup_tenant(test_engine, "SchedDedup")
    TestSessionLocal = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)

    # Insert an overdue obligation
    async with TestSessionLocal() as session:
        contract = Contract(
            organization_id=tenant["org_id"],
            contract_id="CTR-SCHED-01",
            title="Notification Scheduler SLA",
            contract_type="SaaS",
            counterparty="VendorCorp",
            status="ACTIVE"
        )
        session.add(contract)
        await session.flush()

        due = date.today() - timedelta(days=2)
        obligation = Obligation(
            contract_id=contract.id,
            obligor_party="VendorCorp",
            obligee_party="SchedDedup Org",
            title="Overdue Security Assessment Report",
            description="Deliver annual SOC2 audit report",
            due_date=due.isoformat(),
            status="PENDING"
        )
        session.add(obligation)
        await session.commit()

    # Run evaluation cycle #1
    async with TestSessionLocal() as session:
        stats_1 = await NotificationWorker.evaluate_and_dispatch(session, organization_id=tenant["org_id"])
        assert stats_1["notifications_created"] >= 1

    # Run evaluation cycle #2 immediately
    async with TestSessionLocal() as session:
        stats_2 = await NotificationWorker.evaluate_and_dispatch(session, organization_id=tenant["org_id"])
        # Deduplication must suppress second dispatch
        assert stats_2["notifications_created"] == 0
        assert stats_2["skipped_dedup"] >= 1


@pytest.mark.asyncio
async def test_health_check_reports_scheduler_status():
    """Verify the /health endpoint exposes the scheduler status."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.get("/health")
        assert res.status_code == 200
        data = res.json()
        assert data["status"] == "healthy"
        assert "scheduler_running" in data
        assert isinstance(data["scheduler_running"], bool)


# ==============================================================================
# Task B Tests: Downloadable Word (.docx) Redline Export
# ==============================================================================

@pytest.mark.asyncio
async def test_docx_export_success(test_engine):
    """Verify successful generation of a valid DOCX document containing redlines and contract details."""
    tenant = await _setup_tenant(test_engine, "DocxSuccess")
    TestSessionLocal = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)

    async with TestSessionLocal() as session:
        contract = Contract(
            organization_id=tenant["org_id"],
            contract_id="CTR-DOCX-001",
            title="Master Cloud Services Agreement",
            contract_type="MSA",
            counterparty="Enterprise Cloud Inc",
            status="UNDER_REVIEW",
            version="v2.1"
        )
        session.add(contract)
        await session.flush()

        # Add a version with normalized text
        version = ContractVersion(
            contract_id=contract.id,
            version_number="v2.1",
            filename="master_cloud_v2.1.pdf",
            file_storage_path=f"storage/uploads/{tenant['org_id']}/test.pdf",
            file_mime_type="application/pdf",
            file_byte_size=12345,
            sha256_checksum="a" * 64,
            raw_text="The Provider agrees to provide 99.9% uptime availability.",
            normalized_text="The Provider agrees to provide 99.9% uptime availability."
        )
        session.add(version)
        await session.flush()

        # Add redlines: 1 accepted, 1 pending
        r1 = ContractRedline(
            contract_id=contract.id,
            section_title="Service Level Commitments",
            original_text="99.9% uptime availability",
            suggested_text="99.99% uptime availability with quarterly credit audits",
            rationale="Align with enterprise reliability standard",
            status="ACCEPTED",
            author="Legal Reviewer",
            comment="Agreed with vendor during negotiation"
        )
        r2 = ContractRedline(
            contract_id=contract.id,
            section_title="Limitation of Liability",
            original_text="Liability is capped at 1x the annual fee.",
            suggested_text="Liability is capped at 2x annual fees and excludes data security breaches.",
            rationale="Mitigate data protection risk",
            status="PENDING",
            author="Risk Officer",
            comment=None
        )
        session.add_all([r1, r2])
        await session.commit()
        contract_id = contract.id

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.get(
            f"/api/v1/contracts/{contract_id}/export/docx",
            headers=tenant["headers"]
        )

        assert res.status_code == 200
        # Verify Content-Type
        assert res.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        # Verify download headers
        assert "attachment;" in res.headers["content-disposition"]
        assert "CTR-DOCX-001_proposed_revisions.docx" in res.headers["content-disposition"]
        assert "no-store" in res.headers["cache-control"]

        # Verify genuine DOCX binary structure (ZIP magic bytes)
        content_bytes = res.content
        assert content_bytes[:4] == b"PK\x03\x04"

        # Verify ZIP contains Word document elements
        with zipfile.ZipFile(io.BytesIO(content_bytes)) as zf:
            assert "word/document.xml" in zf.namelist()

        # Parse with python-docx and inspect content
        doc = Document(io.BytesIO(content_bytes))
        full_text = "\n".join([p.text for p in doc.paragraphs])
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += f"\n{cell.text}"

        assert "Master Cloud Services Agreement" in full_text
        assert "Proposed Revisions" in full_text
        assert "Service Level Commitments" in full_text
        assert "99.9% uptime availability" in full_text
        assert "99.99% uptime availability" in full_text
        assert "Align with enterprise reliability standard" in full_text

        # Verify audit log was recorded
        async with TestSessionLocal() as audit_session:
            from sqlalchemy.future import select
            audit_stmt = select(AuditLog).where(
                AuditLog.object_id == contract_id,
                AuditLog.object_type == "Contract",
                AuditLog.action == "EXPORT_REDLINES_DOCX"
            )
            audit_res = await audit_session.execute(audit_stmt)
            log = audit_res.scalar_one_or_none()
            assert log is not None
            assert log.user_email == tenant["email"]


@pytest.mark.asyncio
async def test_docx_export_empty_redlines(test_engine):
    """Verify DOCX export succeeds gracefully when a contract has no proposed revisions."""
    tenant = await _setup_tenant(test_engine, "DocxEmpty")
    TestSessionLocal = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)

    async with TestSessionLocal() as session:
        contract = Contract(
            organization_id=tenant["org_id"],
            contract_id="CTR-DOCX-EMPTY",
            title="Clean Non-Disclosure Agreement",
            contract_type="NDA",
            counterparty="PartnerCorp",
            status="DRAFT"
        )
        session.add(contract)
        await session.commit()
        contract_id = contract.id

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.get(
            f"/api/v1/contracts/{contract_id}/export/docx",
            headers=tenant["headers"]
        )

        assert res.status_code == 200
        content_bytes = res.content
        assert content_bytes[:4] == b"PK\x03\x04"

        doc = Document(io.BytesIO(content_bytes))
        full_text = "\n".join([p.text for p in doc.paragraphs])
        assert "Clean Non-Disclosure Agreement" in full_text
        assert "No proposed revisions" in full_text


@pytest.mark.asyncio
async def test_docx_export_not_found(test_engine):
    """Verify 404 is returned when requesting a non-existent contract."""
    tenant = await _setup_tenant(test_engine, "DocxNotFound")
    fake_id = str(uuid.uuid4())

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.get(
            f"/api/v1/contracts/{fake_id}/export/docx",
            headers=tenant["headers"]
        )
        assert res.status_code == 404
        assert "not found" in res.json()["detail"].lower()


@pytest.mark.asyncio
async def test_docx_export_unauthenticated():
    """Verify 401 is returned when accessing the export endpoint without authentication."""
    fake_id = str(uuid.uuid4())
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        res = await client.get(f"/api/v1/contracts/{fake_id}/export/docx")
        assert res.status_code == 401


@pytest.mark.asyncio
async def test_docx_export_cross_tenant_isolation(test_engine):
    """Verify that Tenant A cannot export a contract belonging to Tenant B."""
    tenant_a = await _setup_tenant(test_engine, "DocxTenantA")
    tenant_b = await _setup_tenant(test_engine, "DocxTenantB")
    TestSessionLocal = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)

    async with TestSessionLocal() as session:
        contract_b = Contract(
            organization_id=tenant_b["org_id"],
            contract_id="CTR-TENANT-B-01",
            title="Confidential Tenant B Agreement",
            contract_type="License",
            counterparty="TenantB Partner",
            status="ACTIVE"
        )
        session.add(contract_b)
        await session.commit()
        contract_b_id = contract_b.id

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Tenant A attempts to access Tenant B's contract
        res = await client.get(
            f"/api/v1/contracts/{contract_b_id}/export/docx",
            headers=tenant_a["headers"]
        )
        assert res.status_code == 404


def test_sanitize_filename_utility():
    """Verify filename sanitization removes illegal characters and trims length safely."""
    raw = 'Confidential: "Agreement" & <Annex> / [2026] *?|'
    cleaned = sanitize_filename(raw)
    assert ":" not in cleaned
    assert '"' not in cleaned
    assert "<" not in cleaned
    assert ">" not in cleaned
    assert "/" not in cleaned
    assert "*" not in cleaned
    assert "?" not in cleaned
    assert "|" not in cleaned
    assert cleaned.startswith("Confidential_Agreement")

    # Long name truncation
    long_raw = "A" * 200
    cleaned_long = sanitize_filename(long_raw, max_len=50)
    assert len(cleaned_long) == 50
