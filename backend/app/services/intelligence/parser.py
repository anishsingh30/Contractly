import hashlib
import io
from typing import Dict, Any, List
import pymupdf as fitz
from docx import Document


def compute_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def parse_pdf(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Parses PDF bytes extracting text per page and overall document metadata.
    """
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    page_count = len(doc)
    pages: List[Dict[str, Any]] = []
    full_text_list: List[str] = []

    for page_idx in range(page_count):
        page = doc.load_page(page_idx)
        text = page.get_text("text") or ""
        stripped = text.strip()
        pages.append({
            "page_number": page_idx + 1,
            "text": stripped,
            "char_count": len(stripped)
        })
        if stripped:
            full_text_list.append(stripped)

    doc.close()
    combined_text = "\n\n".join(full_text_list)
    word_count = len(combined_text.split()) if combined_text else 0

    return {
        "raw_text": combined_text,
        "pages": pages,
        "page_count": page_count,
        "char_count": len(combined_text),
        "word_count": word_count,
        "file_byte_size": len(file_bytes),
        "sha256": compute_sha256(file_bytes),
        "mime_type": "application/pdf",
        "filename": filename,
        "parser_status": "COMPLETED"
    }


def parse_docx(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Parses DOCX bytes extracting text from paragraphs and tables.
    """
    doc_stream = io.BytesIO(file_bytes)
    doc = Document(doc_stream)
    full_paragraphs: List[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            full_paragraphs.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                full_paragraphs.append(row_text)

    combined_text = "\n\n".join(full_paragraphs)
    word_count = len(combined_text.split()) if combined_text else 0

    # DOCX does not have fixed physical pages in file format, represents single logical stream
    pages = [{
        "page_number": 1,
        "text": combined_text,
        "char_count": len(combined_text)
    }]

    return {
        "raw_text": combined_text,
        "pages": pages,
        "page_count": 1,
        "char_count": len(combined_text),
        "word_count": word_count,
        "file_byte_size": len(file_bytes),
        "sha256": compute_sha256(file_bytes),
        "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "filename": filename,
        "parser_status": "COMPLETED"
    }


def parse_text_fallback(file_bytes: bytes, filename: str, mime_type: str = "text/plain") -> Dict[str, Any]:
    """
    Fallback parser for plain text, markdown, or unrecognized file byte streams.
    """
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="replace")

    stripped = text.strip()
    word_count = len(stripped.split()) if stripped else 0
    pages = [{
        "page_number": 1,
        "text": stripped,
        "char_count": len(stripped)
    }]

    return {
        "raw_text": stripped,
        "pages": pages,
        "page_count": 1,
        "char_count": len(stripped),
        "word_count": word_count,
        "file_byte_size": len(file_bytes),
        "sha256": compute_sha256(file_bytes),
        "mime_type": mime_type or "text/plain",
        "filename": filename,
        "parser_status": "COMPLETED"
    }


def parse_document(file_bytes: bytes, filename: str, mime_type: str = "") -> Dict[str, Any]:
    """
    Unified entry point for document parsing.
    Dispatches to PDF, DOCX, or text fallback parser.
    """
    lower_name = (filename or "").lower()
    lower_mime = (mime_type or "").lower()

    if lower_name.endswith(".pdf") or "pdf" in lower_mime:
        return parse_pdf(file_bytes, filename)
    elif lower_name.endswith(".docx") or "officedocument" in lower_mime or "word" in lower_mime:
        return parse_docx(file_bytes, filename)
    else:
        return parse_text_fallback(file_bytes, filename, mime_type)
